"""
Reading one source's export into rows the user can review.

Argentine exports write "1.234,56" and "31/12/2026", split money across debit
and credit columns, and carry rows that are not spending at all: transfers
between the user's own accounts, crypto operations, card payments. The Import
Profile describes all of that; this module applies it.
"""

import csv
import io
import re
from dataclasses import dataclass
from datetime import date as Date
from datetime import datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from app.models import ImportProfile
from app.models.enums import Currency, NumberFormat, SignConvention, TransactionType
from app.services.errors import Invalid

NOT_A_NUMBER = re.compile(r"[^\d,.\-]")


@dataclass
class ParsedRow:
    """One line of the export, read but not yet judged."""

    number: int
    date: Date
    description: str
    amount: Decimal
    currency: Currency
    type: TransactionType
    ignored: bool


def read_table(filename: str, content: bytes) -> list[dict[str, str]]:
    """The file as a list of rows keyed by column heading."""
    if filename.lower().endswith(".xlsx"):
        return _read_xlsx(content)
    if filename.lower().endswith(".csv"):
        return _read_csv(content)
    raise Invalid(f"'{filename}' is neither a CSV nor an XLSX export")


def _read_csv(content: bytes) -> list[dict[str, str]]:
    # Exports come out of Argentine tools as UTF-8 or as Latin-1; Latin-1 decodes
    # anything, so it is the fallback rather than a guess that can fail.
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    return [
        {(key or "").strip(): (value or "").strip() for key, value in row.items()}
        for row in reader
    ]


def _read_xlsx(content: bytes) -> list[dict[str, str]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    if sheet is None:
        raise Invalid("the XLSX file has no sheet to read")
    rows = sheet.iter_rows(values_only=True)
    headings = [str(cell).strip() if cell is not None else "" for cell in next(rows, ())]
    return [
        {
            heading: "" if cell is None else str(cell).strip()
            for heading, cell in zip(headings, cells)
        }
        for cells in rows
        if any(cell is not None and str(cell).strip() for cell in cells)
    ]


def read_columns(filename: str, content: bytes) -> list[str]:
    """
    The column headings of an export, so a Profile can be written from a real
    file rather than from memory.
    """
    table = read_table(filename, content)
    if not table:
        raise Invalid(f"'{filename}' has no rows to read column names from")
    return [heading for heading in table[0] if heading]


def parse_number(text: str, number_format: NumberFormat) -> Decimal:
    """"$ 1.234,56" as 1234.56, in whichever way the source writes its numbers."""
    cleaned = NOT_A_NUMBER.sub("", text or "").strip()
    if not cleaned:
        return Decimal(0)
    if number_format == NumberFormat.comma_decimal:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    else:
        cleaned = cleaned.replace(",", "")
    try:
        return Decimal(cleaned)
    except InvalidOperation as error:
        raise Invalid(f"'{text}' is not an amount this Profile can read") from error


def parse_date(text: str, date_format: str) -> Date:
    # Exports often carry a time after the date; the date is all this needs.
    for candidate in (text or "").strip(), (text or "").strip().split(" ")[0]:
        try:
            return datetime.strptime(candidate, date_format).date()
        except ValueError:
            continue
    raise Invalid(f"'{text}' is not a date of the form {date_format}")


def _column(raw: dict[str, str], name: str | None) -> str:
    return "" if name is None else (raw.get(name) or "")


def _amount_and_type(
    raw: dict[str, str], profile: ImportProfile
) -> tuple[Decimal, TransactionType]:
    mapping = profile.column_mapping
    if profile.sign_convention == SignConvention.debit_credit_columns:
        debit = parse_number(_column(raw, mapping.get("debit")), profile.number_format)
        credit = parse_number(
            _column(raw, mapping.get("credit")), profile.number_format
        )
        if debit:
            return abs(debit), TransactionType.expense
        return abs(credit), TransactionType.income

    amount = parse_number(_column(raw, mapping.get("amount")), profile.number_format)
    spends = (
        amount < 0
        if profile.sign_convention == SignConvention.negative_is_expense
        else amount > 0
    )
    return abs(amount), (
        TransactionType.expense if spends else TransactionType.income
    )


def is_ignored(raw: dict[str, str], profile: ImportProfile) -> bool:
    """Rows the Profile says are not spending: transfers, crypto, card payments."""
    mapping = profile.column_mapping
    haystacks = [
        _column(raw, mapping.get("description")).casefold(),
        _column(raw, mapping.get("type")).casefold(),
    ]
    return any(
        pattern.casefold() in haystack
        for pattern in (profile.ignore_patterns or [])
        if pattern
        for haystack in haystacks
    )


def parse_rows(
    table: list[dict[str, str]], profile: ImportProfile
) -> list[ParsedRow]:
    mapping = profile.column_mapping
    parsed = []
    for number, raw in enumerate(table, start=1):
        missing = [
            column
            for column in (mapping["date"], mapping["description"])
            if column not in raw
        ]
        if missing:
            raise Invalid(
                f"the file has no column named {', '.join(repr(c) for c in missing)}"
            )
        amount, type = _amount_and_type(raw, profile)
        currency = _column(raw, mapping.get("currency")).strip().upper()
        parsed.append(
            ParsedRow(
                number=number,
                date=parse_date(_column(raw, mapping["date"]), profile.date_format),
                description=_column(raw, mapping["description"]).strip(),
                amount=amount,
                currency=Currency.USD if currency == "USD" else Currency.ARS,
                type=type,
                ignored=is_ignored(raw, profile),
            )
        )
    return parsed
